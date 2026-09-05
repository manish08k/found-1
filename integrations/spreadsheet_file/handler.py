"""Spreadsheet File integration — read/write CSV and XLSX files."""
import csv
import io
import structlog

from core.execution_engine import register_node

log = structlog.get_logger(__name__)


@register_node("spreadsheet_file.read_csv")
async def read_csv(config: dict, input_data: dict, credential_id: str | None, db) -> dict:
    """Read a CSV file from file_path or parse file_content string.

    config/input_data:
      file_path    — path to CSV file on disk (optional if file_content provided)
      file_content — raw CSV string content (optional if file_path provided)
    """
    merged = {**config, **input_data}
    file_path = merged.get("file_path") or ""
    file_content = merged.get("file_content") or ""

    if file_content:
        reader = csv.DictReader(io.StringIO(file_content))
        rows = list(reader)
        headers = reader.fieldnames or []
    elif file_path:
        with open(file_path, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            headers = reader.fieldnames or []
    else:
        raise ValueError("Either file_path or file_content is required for spreadsheet_file.read_csv")

    log.info("spreadsheet_file.read_csv", rows=len(rows), headers=len(headers))
    return {"rows": rows, "headers": list(headers)}


@register_node("spreadsheet_file.write_csv")
async def write_csv(config: dict, input_data: dict, credential_id: str | None, db) -> dict:
    """Write rows to a CSV file.

    config:
      file_path — destination file path (required)
    input_data:
      rows      — list of dicts to write (required)
    """
    merged = {**config, **input_data}
    file_path = merged.get("file_path") or ""
    rows = merged.get("rows") or []

    if not file_path:
        raise ValueError("file_path is required for spreadsheet_file.write_csv")
    if not rows:
        raise ValueError("rows is required for spreadsheet_file.write_csv")

    headers = list(rows[0].keys()) if rows else []

    with open(file_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)

    log.info("spreadsheet_file.write_csv", written=len(rows), file_path=file_path)
    return {"written": len(rows), "file_path": file_path}


@register_node("spreadsheet_file.read_xlsx")
async def read_xlsx(config: dict, input_data: dict, credential_id: str | None, db) -> dict:
    """Read an XLSX file. Requires openpyxl.

    config/input_data:
      file_path  — path to XLSX file (required)
      sheet_name — sheet to read (optional, defaults to active sheet)
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "openpyxl is required for spreadsheet_file.read_xlsx. "
            "Install it with: pip install openpyxl"
        )

    merged = {**config, **input_data}
    file_path = merged.get("file_path") or ""
    sheet_name = merged.get("sheet_name") or None

    if not file_path:
        raise ValueError("file_path is required for spreadsheet_file.read_xlsx")

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    actual_sheet = ws.title

    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not all_rows:
        return {"rows": [], "sheet": actual_sheet}

    headers = [str(h) if h is not None else "" for h in all_rows[0]]
    rows = []
    for raw_row in all_rows[1:]:
        rows.append({headers[i]: raw_row[i] for i in range(len(headers))})

    log.info("spreadsheet_file.read_xlsx", rows=len(rows), sheet=actual_sheet)
    return {"rows": rows, "sheet": actual_sheet}


@register_node("spreadsheet_file.write_xlsx")
async def write_xlsx(config: dict, input_data: dict, credential_id: str | None, db) -> dict:
    """Write rows to an XLSX file. Requires openpyxl.

    config:
      file_path  — destination file path (required)
      sheet_name — sheet name (optional, default 'Sheet1')
    input_data:
      rows       — list of dicts to write (required)
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "openpyxl is required for spreadsheet_file.write_xlsx. "
            "Install it with: pip install openpyxl"
        )

    merged = {**config, **input_data}
    file_path = merged.get("file_path") or ""
    rows = merged.get("rows") or []
    sheet_name = merged.get("sheet_name") or "Sheet1"

    if not file_path:
        raise ValueError("file_path is required for spreadsheet_file.write_xlsx")
    if not rows:
        raise ValueError("rows is required for spreadsheet_file.write_xlsx")

    headers = list(rows[0].keys()) if rows else []

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])

    wb.save(file_path)

    log.info("spreadsheet_file.write_xlsx", written=len(rows), file_path=file_path)
    return {"written": len(rows)}
